"""Tables workbook with Total / By banner / Stat sheets (TabulateSpec.output_mode = "banner_sheets").

Sheets:
  Total              every table, Total column only
  By <banner>        every table by one banner group, with a Total column last
  By <banner> (Stat) the same with a row of column-comparison letters under each value

Each table: title, header ("Column %" + columns; a grid gets one column group per item),
answer rows, nets, NET (anyone answering), Column Sample Size, then the title again and
"sample size = …". Values are fractions formatted as percentages, so the cells stay numeric.

Conventions:
  - nets that split all of a question's answers (1 + 2 / 3 + 4 + 5 / 6 + 7) are a merge:
    they replace the answers instead of being added under them
  - rows nobody gave (Total 0 across the table) are left out
  - letters: uppercase at the higher confidence level, lowercase at the lower one; the
    Total column isn't tested ("-"), nor a column nobody in the table answered
  - hide_codes: answers (rows) and banner values (columns) left out of the tables; bases and
    the other columns' tests are unchanged, and letters pointing at a hidden column are dropped
  - show_row_pct: one more sheet per banner, "(Row %)": each column's share of the row's Total
"""

import io
import re
from dataclasses import dataclass, field
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services.data_loader import is_non_substantive
from shared.significance import significant_at, z_test_proportions

TITLE = Font(size=11)
SHEET_TITLE = Font(size=11, bold=True)
HEAD = Font(size=10, bold=True)
LABEL = Font(size=10, bold=True)
VALUE = Font(size=10)
WHITE = PatternFill("solid", fgColor="FFFFFF")
BASE_FILL = PatternFill("solid", fgColor="E8EBF6")
BASE_BORDER = Border(top=Side(style="medium"), bottom=Side(style="medium"))
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left")
RIGHT = Alignment(horizontal="right")
GAP = 3  # blank rows between tables
FIRST_ROW = 5


@dataclass
class Row:
    key: Any                 # answer code / MRS member, ("net", label), or "NET"
    label: str
    kind: str                # "answer", "net" or "NET"
    codes: list = field(default_factory=list)   # nets: the answers added up


@dataclass
class Item:
    """One question in a table (a grid table has several)."""
    label: str
    sheet: Any               # SheetResult


@dataclass
class Table:
    key: str                 # stub variable or group name (TabulateSpec.table_order / table_titles)
    title: str
    items: list[Item]
    rows: list[Row] = field(default_factory=list)


# ── build: SheetResults → tables ─────────────────────────────────────────────

def _tables(result, spec) -> list[Table]:
    tables, grids = [], {}
    for sheet in result.sheets:
        if sheet.status != "success" or not sheet.crosstab_data:
            continue
        if sheet.group:
            if sheet.group not in grids:
                grids[sheet.group] = Table(key=sheet.group, title=sheet.group, items=[])
                tables.append(grids[sheet.group])
            grids[sheet.group].items.append(Item(label=sheet.label or sheet.variable, sheet=sheet))
        else:
            tables.append(Table(key=sheet.variable, title=sheet.label or sheet.variable, items=[Item(label="", sheet=sheet)]))
    for table in grids.values():  # items whose labels were cut to the same text: say which variable
        labels = [i.label for i in table.items]
        for item in table.items:
            if labels.count(item.label) > 1:
                item.label = f"{item.label} ({item.sheet.variable})"
    titles = spec.table_titles or {}
    for table in tables:
        table.title = titles.get(table.key) or table.title
    if spec.table_order:
        rank = {key: i for i, key in enumerate(spec.table_order)}
        tables.sort(key=lambda t: rank.get(t.key, len(rank)))  # stable: unlisted tables keep their order
    return tables


def _nets(spec, sheet) -> dict[str, list]:
    return {} if sheet.is_mrs else dict((spec.nets or {}).get(sheet.variable, {}))


def _is_merge(nets: dict[str, list], answer_codes: set) -> bool:
    """Nets that don't overlap and cover every answer given replace the answers."""
    if not nets or not answer_codes:
        return False
    seen = [float(c) for codes in nets.values() for c in codes]
    return len(seen) == len(set(seen)) and answer_codes <= set(seen)


def _answer_rows(sheet) -> list[Row]:
    rows, seen = [], set()
    for ct in sheet.crosstab_data.values():
        for rd in ct.get("table", []):
            rv = rd.get("row_value")
            if rv not in seen:
                seen.add(rv)
                rows.append(Row(key=rv, label=str(rd.get("row_label", rv)), kind="answer"))
    if not sheet.is_mrs:
        rows.sort(key=lambda r: (not isinstance(r.key, (int, float)), r.key if isinstance(r.key, (int, float)) else 0))
    return rows


def _code(x):
    """Codes compare as numbers when they are numbers (1, 1.0 and "1" are the same answer)."""
    try:
        return float(x)
    except (TypeError, ValueError):
        return str(x)


def _hidden(spec, *names) -> set:
    hide = spec.hide_codes or {}
    return {_code(c) for name in names for c in hide.get(name, [])}


def _layout_rows(table: Table, spec, calc) -> None:
    """Rows of a table: answers (or the nets that merge them), other nets, NET; empty and hidden rows out."""
    rows: list[Row] = []
    for item in table.items:
        answers = _answer_rows(item.sheet)
        nets = _nets(spec, item.sheet)
        merge = _is_merge(nets, {float(r.key) for r in answers if isinstance(r.key, (int, float))})
        candidates = ([] if merge else answers) + [Row(key=("net", k), label=k, kind="net", codes=list(v))
                                                  for k, v in nets.items()]
        for r in candidates:
            if all(r.key != x.key for x in rows):
                rows.append(r)
    # A row nobody gave (Total) in any item is left out
    rows = [r for r in rows if any(calc.count(item, calc.total, r) > 0 for item in table.items)]
    hidden = _hidden(spec, table.key, *(item.sheet.variable for item in table.items))
    rows = [r for r in rows if not (r.kind == "answer" and _code(r.key) in hidden)]
    if spec.sort_rows:
        answers = [r for r in rows if r.kind == "answer"]
        last = [r for r in answers if _goes_last(r.label)]
        firsts = [r for r in answers if r not in last]
        if spec.sort_rows == "az":
            firsts.sort(key=lambda r: r.label.casefold())
        else:
            firsts.sort(key=lambda r: -max(calc.value(item, calc.total, r) or 0 for item in table.items))
        rows = firsts + last + [r for r in rows if r.kind == "net"]
    table.rows = rows + [Row(key="NET", label="NET", kind="NET")]


def _goes_last(label: str) -> bool:
    low = label.strip().lower()
    return is_non_substantive(label) or low.startswith(("other", "none", "no ", "nothing"))


class Calc:
    """Counts, bases, percentages and letters per item × column × row."""

    def __init__(self, result, spec, data):
        self.spec, self.df = spec, data.df
        self.total = next((bc for bc in result.banner_columns if bc.banner_var == "_total_"), None)

    @staticmethod
    def _ct(item, bc) -> dict:
        return (item.sheet.crosstab_data or {}).get(bc.banner_var, {}) if bc else {}

    def _cell(self, item, bc, key) -> dict | None:
        for rd in self._ct(item, bc).get("table", []):
            if rd.get("row_value") == key:
                cell = rd.get(bc.value)
                return cell if isinstance(cell, dict) else None
        return None

    def base(self, item, bc) -> float:
        ct = self._ct(item, bc)
        if ct.get("is_mrs"):
            return float((ct.get("col_bases") or {}).get(bc.value, 0))
        return float(sum(rd[bc.value].get("count", 0) for rd in ct.get("table", [])
                         if isinstance(rd.get(bc.value), dict)))

    def count(self, item, bc, row: Row) -> float:
        if bc is None:
            return 0.0
        if row.kind == "answer":
            cell = self._cell(item, bc, row.key)
            return float(cell.get("count", 0)) if cell else 0.0
        if row.kind == "net":
            wanted = {float(c) for c in row.codes}
            return float(sum(rd[bc.value].get("count", 0) for rd in self._ct(item, bc).get("table", [])
                             if isinstance(rd.get(bc.value), dict) and isinstance(rd.get("row_value"), (int, float))
                             and float(rd["row_value"]) in wanted))
        return self._net_count(item, bc)

    def _net_count(self, item, bc) -> float:
        """NET: respondents in the base who gave any answer (a multi-select: ticked any option)."""
        sheet = item.sheet
        if not sheet.is_mrs:
            return self.base(item, bc)
        members = [m for m in (sheet.mrs_members or []) if m in self.df.columns]
        col = self.df.get(bc.banner_var)
        if col is None or not members:
            return 0.0
        in_col = col == _value(bc.value)
        answered = self.df[members].notna().any(axis=1)
        ticked = (self.df[members].fillna(0) > 0).any(axis=1)
        selected = in_col & answered & ticked
        if self.spec.weight and self.spec.weight in self.df.columns:
            return float(self.df.loc[selected, self.spec.weight].sum())
        return float(selected.sum())

    def value(self, item, bc, row: Row) -> float | None:
        base = self.base(item, bc) if bc else 0
        return self.count(item, bc, row) / base if base else None

    def row_value(self, item, bc, row: Row) -> float | None:
        """Row %: this column's share of the row's Total."""
        total = self.count(item, self.total, row) if self.total else 0
        return self.count(item, bc, row) / total if total else None

    def letters(self, item, cols, bc, row: Row) -> str:
        """Columns this one is significantly higher than ('-' where there is no test)."""
        if bc.banner_var == "_total_" or row.kind == "NET" or not self.base(item, bc):
            return "-"
        if sum(1 for c in cols if self.base(item, c)) < 2:
            return "-"
        sheet_letter = {c.value: get_column_letter(i + 1) for i, c in enumerate(cols)}
        if row.kind == "answer":
            cell = self._cell(item, bc, row.key) or {}
            ct = self._ct(item, bc)
            value_of = {letter: value for value, letter in ct.get("col_labels", {}).items()}
            hi = [sheet_letter.get(value_of.get(x)) for x in cell.get("significance_letters", [])]
            lo = [sheet_letter.get(value_of.get(x)) for x in cell.get("significance_letters_secondary", [])]
        else:
            hi, lo = self._test_net(item, cols, bc, row)
        return " ".join([x for x in hi if x] + [x.lower() for x in lo if x])

    def _test_net(self, item, cols, bc, row):
        alpha = 1 - self.spec.significance_level
        alpha2 = self.spec.secondary_alpha
        p1, n1 = self.value(item, bc, row) or 0, self.base(item, bc)
        hi, lo = [], []
        for i, other in enumerate(cols):
            if other is bc or not self.base(item, other):
                continue
            p2, n2 = self.value(item, other, row) or 0, self.base(item, other)
            p, sig = z_test_proportions(p1, n1, p2, n2, alpha, len(cols))
            if sig:
                hi.append(get_column_letter(i + 1))
            elif significant_at(p, alpha2, len(cols), p1 > p2):
                lo.append(get_column_letter(i + 1))
        return hi, lo


def _value(v: str):
    try:
        return float(v)
    except ValueError:
        return v


# ── render ───────────────────────────────────────────────────────────────────

def _put(ws, row, col, value, font=VALUE, align=RIGHT, fmt=None, fill=WHITE, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font, cell.alignment = font, align
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border
    return cell


def _sample_line(calc, table, n_cases) -> str:
    n = int(max((calc.base(item, calc.total) for item in table.items), default=0))
    return f"sample size = {n}" + (f"; total sample size = {n_cases}; {n_cases - n} missing" if n < n_cases else "")


def _write_table(ws, r, table: Table, calc, cols, stat: bool, banner: str | None, n_cases: int, row_pct: bool = False) -> int:
    """One table at row r; returns the next free row. row_pct: values are Row % (no letters)."""
    value_of = calc.row_value if row_pct else calc.value
    title = f"{table.title} by {banner}" if banner else table.title
    _put(ws, r, 1, title, font=TITLE, align=None, fill=None)
    r += 1
    grid = len(table.items) > 1
    corner = ("Row %" if row_pct else "Column %") + ("\nColumn Comparisons" if stat else "")
    if not banner and not grid:
        corner = " "
    width = len(cols)
    if grid:
        _put(ws, r, 1, corner, font=HEAD, align=CENTER)
        _put(ws, r + 1, 1, None, font=HEAD, align=CENTER)
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
        for k, item in enumerate(table.items):
            c0 = 2 + k * width
            _put(ws, r, c0, item.label, font=HEAD, align=CENTER)
            if width > 1:
                ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + width - 1)
            for j, (label, _) in enumerate(cols):
                _put(ws, r + 1, c0 + j, label, font=HEAD, align=CENTER)
        r += 2
    else:
        _put(ws, r, 1, corner, font=HEAD, align=CENTER)
        for j, (label, _) in enumerate(cols):
            _put(ws, r, 2 + j, label if banner else "%", font=HEAD, align=CENTER)
        r += 1

    tested = [bc for _, bc in cols if bc.banner_var != "_total_"]
    names = 1 if stat and grid else 0  # a grid repeats the column letters under each value
    for row in table.rows:
        _put(ws, r, 1, row.label, font=LABEL, align=LEFT)
        for k, item in enumerate(table.items):
            for j, (_, bc) in enumerate(cols):
                v = value_of(item, bc, row)
                _put(ws, r, 2 + k * width + j, v, fmt="0%")
                if names:
                    _put(ws, r + 1, 2 + k * width + j, get_column_letter(j + 1), align=CENTER)
                if stat:
                    _put(ws, r + 1 + names, 2 + k * width + j, calc.letters(item, tested, bc, row) if v is not None else "-")
        if stat:
            for extra in range(1 + names):
                _put(ws, r + 1 + extra, 1, None, font=LABEL, align=LEFT)
        r += 1 + (1 + names if stat else 0)

    if banner and not grid:
        _put(ws, r, 1, "Column Sample Size", font=LABEL, align=CENTER, fill=BASE_FILL, border=BASE_BORDER)
        for j, (_, bc) in enumerate(cols):
            _put(ws, r, 2 + j, int(round(calc.base(table.items[0], bc))), fmt="0", fill=BASE_FILL, border=BASE_BORDER)
        r += 1
        if stat:
            _put(ws, r, 1, "Column Names", font=LABEL, align=LEFT)
            for j in range(width):
                _put(ws, r, 2 + j, get_column_letter(j + 1), align=CENTER)
            r += 1
    _put(ws, r, 1, title if banner else f"{table.title} SUMMARY", font=TITLE, align=None, fill=None)
    footer = _sample_line(calc, table, n_cases)
    if stat:
        footer += "; " + _stat_note(calc.spec)
    _put(ws, r + 1, 1, footer, font=TITLE, align=None, fill=None)
    return r + 2 + GAP


def _pct(level: float) -> str:
    return f"{level * 100:.1f}".rstrip("0").rstrip(".") + "%"


def _stat_note(spec) -> str:
    if spec.secondary_level is None:
        symbols = f"A, B, C... (confidence level >= {_pct(spec.significance_level)})"
    else:
        symbols = (f"a, b, c... (confidence level >= {_pct(spec.secondary_level)}), "
                   f"A, B, C... (confidence level >= {_pct(spec.significance_level)})")
    note = (f"{_pct(spec.secondary_level or spec.significance_level)} confidence level; "
            f"Column comparison symbols: {symbols}; No test symbol: -")
    return note + ("" if spec.bonferroni else "; no multiple-comparison correction")


def _sheet(wb, title: str):
    safe = re.sub(r"[\\/*?:\[\]]", "_", title).strip("'") or "Tables"
    name, suffix = safe[:31], 1
    while name.casefold() in {s.casefold() for s in wb.sheetnames}:
        suffix += 1
        ending = f" ({suffix})"
        name = safe[:31 - len(ending)] + ending
    ws = wb.create_sheet(title=name)
    _put(ws, 1, 1, title, font=SHEET_TITLE, align=None, fill=None)
    ws.column_dimensions["A"].width = 60
    return ws


def _widths(ws, n_cols: int):
    from openpyxl.utils import get_column_letter
    for c in range(2, n_cols + 2):
        ws.column_dimensions[get_column_letter(c)].width = 13


def _banner_groups(result, spec) -> list[list]:
    """Banner columns by banner group (the Total column on its own), hidden values left out."""
    groups: dict[int, list] = {}
    for bc in result.banner_columns:
        if bc.banner_var != "_total_" and _code(bc.value) not in _hidden(spec, bc.banner_var):
            groups.setdefault(bc.banner_index, []).append(bc)
    return list(groups.values())


@dataclass
class Placed:
    """A laid-out table and what it was tabulated with: its calculations and its column groups."""
    table: Table
    calc: Calc
    groups: list


def _placed(result, spec, data) -> list[Placed]:
    """Every table in output order. With per-table columns (result.parts) each table keeps its own part."""
    parts = [(p.result, p.spec, p.data) for p in result.parts] or [(result, spec, data)]
    out = []
    for r, s, d in parts:
        calc, groups = Calc(r, s, d), _banner_groups(r, s)
        for table in _tables(r, s):
            _layout_rows(table, s, calc)
            out.append(Placed(table, calc, groups))
    if result.parts and spec.table_order:
        rank = {key: i for i, key in enumerate(spec.table_order)}
        out.sort(key=lambda x: rank.get(x.table.key, len(rank)))
    return out


def _group_keys(groups) -> list[tuple]:
    """(banner label, nth group with that label): custom groups that overlap share a label."""
    seen, keys = {}, []
    for g in groups:
        label = g[0].banner_label
        keys.append((label, seen.get(label, 0)))
        seen[label] = seen.get(label, 0) + 1
    return keys


def _banners_json(groups) -> list[dict]:
    return [{"variable": g[0].banner_var, "label": g[0].banner_label,
             "columns": [{"value": bc.value, "label": bc.value_label, "letter": get_column_letter(i + 1)}
                         for i, bc in enumerate(g)]} for g in groups]


def preview_tables(result, spec, data) -> dict:
    """The tables this workbook writes, as JSON for a live preview: the same rows (answers or
    merging nets, nets, NET; empty rows out; sorted if asked), bases, percentages and the Stat
    sheets' letters (each banner group tested on its own, columns lettered A, B, … within it).
    Each table lists its own banners (per-table columns); the top-level banners are the first
    table's. Call inside multiple_comparison_correction(spec.bonferroni), like the workbook."""
    placed = _placed(result, spec, data)

    def cell(calc, item, bc, tested, row) -> dict:
        v = calc.value(item, bc, row)
        letters = calc.letters(item, tested, bc, row) if v is not None and tested else "-"
        return {"count": round(calc.count(item, bc, row), 4), "pct": v, "row_pct": calc.row_value(item, bc, row), "sig": letters}

    def item_json(p, item) -> dict:
        calc, total, rows = p.calc, p.calc.total, p.table.rows
        return {
            "label": item.label,
            "variable": item.sheet.variable,
            "is_mrs": item.sheet.is_mrs,
            "total": {"base": calc.base(item, total), "cells": [cell(calc, item, total, [], r) for r in rows]} if total else None,
            "banners": [{"bases": [calc.base(item, bc) for bc in group],
                         "cells": [[cell(calc, item, bc, group, r) for bc in group] for r in rows]} for group in p.groups],
        }

    first_groups = placed[0].groups if placed else _banner_groups(result, spec)
    return {
        "n_cases": len(data.df),
        "banners": _banners_json(first_groups),
        "tables": [{"key": p.table.key, "title": p.table.title, "is_grid": len(p.table.items) > 1,
                    "banners": _banners_json(p.groups),
                    "rows": [{"label": r.label, "kind": r.kind} for r in p.table.rows],
                    "items": [item_json(p, item) for item in p.table.items]} for p in placed],
        "failed": [{"variable": s.variable, "error": s.error} for s in result.sheets if s.status != "success"],
        "significance": {"level": spec.significance_level, "secondary_level": spec.secondary_level,
                         "bonferroni": spec.bonferroni, "note": _stat_note(spec)},
    }


def build_banner_workbook(result, spec, data) -> bytes:
    placed = _placed(result, spec, data)
    n_cases = len(data.df)
    width = lambda p: max(len(p.table.items), 1)

    wb = Workbook()
    wb.remove(wb.active)
    with_total = [p for p in placed if p.calc.total]
    if with_total:
        ws = _sheet(wb, "Total")
        r = FIRST_ROW
        for p in with_total:
            r = _write_table(ws, r, p.table, p.calc, [("Total", p.calc.total)], stat=False, banner=None, n_cases=n_cases)
        _widths(ws, max((width(p) for p in with_total), default=1))

    # One set of sheets per banner; each holds the tables that use that banner (all of them, unless
    # tables have their own columns)
    keys = []
    for p in placed:
        for key in _group_keys(p.groups):
            if key not in keys:
                keys.append(key)
    sheets = [("", False, False), (" (Stat)", True, False)] + ([(" (Row %)", False, True)] if spec.show_row_pct else [])
    for key in keys:
        banner = key[0]
        for suffix, stat, row_pct in sheets:
            ws = _sheet(wb, f"By {banner}{suffix}")
            r, widest = FIRST_ROW, 1
            for p in placed:
                group = dict(zip(_group_keys(p.groups), p.groups)).get(key)
                if group is None:
                    continue
                cols = [(bc.value_label, bc) for bc in group] + ([("Total", p.calc.total)] if p.calc.total else [])
                r = _write_table(ws, r, p.table, p.calc, cols, stat=stat, banner=banner, n_cases=n_cases, row_pct=row_pct)
                widest = max(widest, len(cols) * width(p))
            _widths(ws, widest)

    if not wb.worksheets:
        wb.create_sheet("Tables")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
