"""Tables workbook laid out like a Displayr / Q export (TabulateSpec.output_mode = "displayr").

Sheets:
  Total              every table, Total column only
  By <banner>        every table by one banner group, with a Total column last
  By <banner> (Stat) the same with a row of column-comparison letters under each value

Each table: title, header ("Column %" + columns; a grid gets one column group per item),
answer rows, nets, NET (anyone answering), Column Sample Size, then the title again and
"sample size = …". Values are fractions formatted as percentages, so the cells stay numeric.

Displayr conventions followed:
  - nets that split all of a question's answers (1 + 2 / 3 + 4 + 5 / 6 + 7) are a merge:
    they replace the answers instead of being added under them
  - rows nobody gave (Total 0 across the table) are left out
  - letters: uppercase at the higher confidence level, lowercase at the lower one; the
    Total column isn't tested ("-"), nor a column nobody in the table answered
"""

import io
import string
from dataclasses import dataclass, field
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from services.data_loader import is_non_substantive
from shared.significance import significant_at, z_test_proportions

TITLE = Font(size=11)
SHEET_TITLE = Font(size=11, bold=True)
HEAD = Font(size=10, bold=True)
LABEL = Font(size=10, bold=True)
VALUE = Font(size=10)
WHITE = PatternFill("solid", fgColor="FFFFFF")
BASE_FILL = PatternFill("solid", fgColor="E8EBF6")
BASE_BORDER = Border(top=Side(style="medium"), bottom=Side(style="medium"))
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left")
RIGHT = Alignment(horizontal="right")
GAP = 3  # blank rows between tables
FIRST_ROW = 5


@dataclass
class Row:
    key: Any                 # answer code / MRS member, ("net", label), or "NET"
    label: str
    kind: str                # "answer", "net" or "NET"
    codes: list = field(default_factory=list)   # nets: the answers added up


@dataclass
class Item:
    """One question in a table (a grid table has several)."""
    label: str
    sheet: Any               # SheetResult


@dataclass
class Table:
    key: str                 # stub variable or group name (TabulateSpec.table_order / table_titles)
    title: str
    items: list[Item]
    rows: list[Row] = field(default_factory=list)


# ── build: SheetResults → tables ─────────────────────────────────────────────

def _tables(result, spec) -> list[Table]:
    tables, grids = [], {}
    for sheet in result.sheets:
        if sheet.status != "success" or not sheet.crosstab_data:
            continue
        if sheet.group:
            if sheet.group not in grids:
                grids[sheet.group] = Table(key=sheet.group, title=sheet.group, items=[])
                tables.append(grids[sheet.group])
            grids[sheet.group].items.append(Item(label=sheet.label or sheet.variable, sheet=sheet))
        else:
            tables.append(Table(key=sheet.variable, title=sheet.label or sheet.variable, items=[Item(label="", sheet=sheet)]))
    for table in grids.values():  # items whose labels were cut to the same text: say which variable
        labels = [i.label for i in table.items]
        for item in table.items:
            if labels.count(item.label) > 1:
                item.label = f"{item.label} ({item.sheet.variable})"
    titles = spec.table_titles or {}
    for table in tables:
        table.title = titles.get(table.key) or table.title
    if spec.table_order:
        rank = {key: i for i, key in enumerate(spec.table_order)}
        tables.sort(key=lambda t: rank.get(t.key, len(rank)))  # stable: unlisted tables keep their order
    return tables


def _nets(spec, sheet) -> dict[str, list]:
    return {} if sheet.is_mrs else dict((spec.nets or {}).get(sheet.variable, {}))


def _is_merge(nets: dict[str, list], answer_codes: set) -> bool:
    """Nets that don't overlap and cover every answer given replace the answers (Displayr's merge)."""
    if not nets or not answer_codes:
        return False
    seen = [float(c) for codes in nets.values() for c in codes]
    return len(seen) == len(set(seen)) and answer_codes <= set(seen)


def _answer_rows(sheet) -> list[Row]:
    rows, seen = [], set()
    for ct in sheet.crosstab_data.values():
        for rd in ct.get("table", []):
            rv = rd.get("row_value")
            if rv not in seen:
                seen.add(rv)
                rows.append(Row(key=rv, label=str(rd.get("row_label", rv)), kind="answer"))
    if not sheet.is_mrs:
        rows.sort(key=lambda r: (not isinstance(r.key, (int, float)), r.key if isinstance(r.key, (int, float)) else 0))
    return rows


def _layout_rows(table: Table, spec, calc) -> None:
    """Rows of a table: answers (or the nets that merge them), other nets, NET; empty rows out."""
    rows: list[Row] = []
    for item in table.items:
        answers = _answer_rows(item.sheet)
        nets = _nets(spec, item.sheet)
        merge = _is_merge(nets, {float(r.key) for r in answers if isinstance(r.key, (int, float))})
        candidates = ([] if merge else answers) + [Row(key=("net", k), label=k, kind="net", codes=list(v))
                                                  for k, v in nets.items()]
        for r in candidates:
            if all(r.key != x.key for x in rows):
                rows.append(r)
    # A row nobody gave (Total) in any item is left out, as Displayr does
    rows = [r for r in rows if any(calc.count(item, calc.total, r) > 0 for item in table.items)]
    if spec.sort_rows:
        answers = [r for r in rows if r.kind == "answer"]
        last = [r for r in answers if _goes_last(r.label)]
        firsts = [r for r in answers if r not in last]
        firsts.sort(key=lambda r: -max(calc.value(item, calc.total, r) or 0 for item in table.items))
        rows = firsts + last + [r for r in rows if r.kind == "net"]
    table.rows = rows + [Row(key="NET", label="NET", kind="NET")]


def _goes_last(label: str) -> bool:
    low = label.strip().lower()
    return is_non_substantive(label) or low.startswith(("other", "none", "no ", "nothing"))


class Calc:
    """Counts, bases, percentages and letters per item × column × row."""

    def __init__(self, result, spec, data):
        self.spec, self.df = spec, data.df
        self.total = next((bc for bc in result.banner_columns if bc.banner_var == "_total_"), None)

    @staticmethod
    def _ct(item, bc) -> dict:
        return (item.sheet.crosstab_data or {}).get(bc.banner_var, {}) if bc else {}

    def _cell(self, item, bc, key) -> dict | None:
        for rd in self._ct(item, bc).get("table", []):
            if rd.get("row_value") == key:
                cell = rd.get(bc.value)
                return cell if isinstance(cell, dict) else None
        return None

    def base(self, item, bc) -> float:
        ct = self._ct(item, bc)
        if ct.get("is_mrs"):
            return float((ct.get("col_bases") or {}).get(bc.value, 0))
        return float(sum(rd[bc.value].get("count", 0) for rd in ct.get("table", [])
                         if isinstance(rd.get(bc.value), dict)))

    def count(self, item, bc, row: Row) -> float:
        if bc is None:
            return 0.0
        if row.kind == "answer":
            cell = self._cell(item, bc, row.key)
            return float(cell.get("count", 0)) if cell else 0.0
        if row.kind == "net":
            wanted = {float(c) for c in row.codes}
            return float(sum(rd[bc.value].get("count", 0) for rd in self._ct(item, bc).get("table", [])
                             if isinstance(rd.get(bc.value), dict) and isinstance(rd.get("row_value"), (int, float))
                             and float(rd["row_value"]) in wanted))
        return self._net_count(item, bc)

    def _net_count(self, item, bc) -> float:
        """NET: respondents in the base who gave any answer (a multi-select: ticked any option)."""
        sheet = item.sheet
        if not sheet.is_mrs:
            return self.base(item, bc)
        members = [m for m in (sheet.mrs_members or []) if m in self.df.columns]
        col = self.df.get(bc.banner_var)
        if col is None or not members:
            return 0.0
        in_col = col == _value(bc.value)
        answered = self.df[members].notna().any(axis=1)
        ticked = (self.df[members].fillna(0) > 0).any(axis=1)
        return float((in_col & answered & ticked).sum())

    def value(self, item, bc, row: Row) -> float | None:
        base = self.base(item, bc) if bc else 0
        return self.count(item, bc, row) / base if base else None

    def letters(self, item, cols, bc, row: Row) -> str:
        """Columns this one is significantly higher than ('-' where there is no test)."""
        if bc.banner_var == "_total_" or row.kind == "NET" or not self.base(item, bc):
            return "-"
        if sum(1 for c in cols if self.base(item, c)) < 2:
            return "-"
        sheet_letter = {c.value: string.ascii_uppercase[i] for i, c in enumerate(cols)}
        if row.kind == "answer":
            cell = self._cell(item, bc, row.key) or {}
            ct = self._ct(item, bc)
            value_of = {letter: value for value, letter in ct.get("col_labels", {}).items()}
            hi = [sheet_letter.get(value_of.get(x)) for x in cell.get("significance_letters", [])]
            lo = [sheet_letter.get(value_of.get(x)) for x in cell.get("significance_letters_secondary", [])]
        else:
            hi, lo = self._test_net(item, cols, bc, row)
        return " ".join([x for x in hi if x] + [x.lower() for x in lo if x])

    def _test_net(self, item, cols, bc, row):
        alpha = 1 - self.spec.significance_level
        alpha2 = self.spec.secondary_alpha
        p1, n1 = self.value(item, bc, row) or 0, self.base(item, bc)
        hi, lo = [], []
        for i, other in enumerate(cols):
            if other is bc or not self.base(item, other):
                continue
            p2, n2 = self.value(item, other, row) or 0, self.base(item, other)
            p, sig = z_test_proportions(p1, n1, p2, n2, alpha, len(cols))
            if sig:
                hi.append(string.ascii_uppercase[i])
            elif significant_at(p, alpha2, len(cols), p1 > p2):
                lo.append(string.ascii_uppercase[i])
        return hi, lo


def _value(v: str):
    try:
        return float(v)
    except ValueError:
        return v


# ── render ───────────────────────────────────────────────────────────────────

def _put(ws, row, col, value, font=VALUE, align=RIGHT, fmt=None, fill=WHITE, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font, cell.alignment = font, align
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border
    return cell


def _sample_line(calc, table, n_cases) -> str:
    n = int(max((calc.base(item, calc.total) for item in table.items), default=0))
    return f"sample size = {n}" + (f"; total sample size = {n_cases}; {n_cases - n} missing" if n < n_cases else "")


def _write_table(ws, r, table: Table, calc, cols, stat: bool, banner: str | None, n_cases: int) -> int:
    """One table at row r; returns the next free row."""
    title = f"{table.title} by {banner}" if banner else table.title
    _put(ws, r, 1, title, font=TITLE, align=None, fill=None)
    r += 1
    grid = len(table.items) > 1
    corner = "Column %" + ("\nColumn Comparisons" if stat else "")
    if not banner and not grid:
        corner = " "
    width = len(cols)
    if grid:
        _put(ws, r, 1, corner, font=HEAD, align=CENTER)
        _put(ws, r + 1, 1, None, font=HEAD, align=CENTER)
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
        for k, item in enumerate(table.items):
            c0 = 2 + k * width
            _put(ws, r, c0, item.label, font=HEAD, align=CENTER)
            if width > 1:
                ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + width - 1)
            for j, (label, _) in enumerate(cols):
                _put(ws, r + 1, c0 + j, label, font=HEAD, align=CENTER)
        r += 2
    else:
        _put(ws, r, 1, corner, font=HEAD, align=CENTER)
        for j, (label, _) in enumerate(cols):
            _put(ws, r, 2 + j, label if banner else "%", font=HEAD, align=CENTER)
        r += 1

    tested = [bc for _, bc in cols if bc.banner_var != "_total_"]
    names = 1 if stat and grid else 0  # a grid repeats the column letters under each value
    for row in table.rows:
        _put(ws, r, 1, row.label, font=LABEL, align=LEFT)
        for k, item in enumerate(table.items):
            for j, (_, bc) in enumerate(cols):
                v = calc.value(item, bc, row)
                _put(ws, r, 2 + k * width + j, v, fmt="0%")
                if names:
                    _put(ws, r + 1, 2 + k * width + j, string.ascii_uppercase[j], align=CENTER)
                if stat:
                    _put(ws, r + 1 + names, 2 + k * width + j, calc.letters(item, tested, bc, row) if v is not None else "-")
        if stat:
            for extra in range(1 + names):
                _put(ws, r + 1 + extra, 1, None, font=LABEL, align=LEFT)
        r += 1 + (1 + names if stat else 0)

    if banner and not grid:
        _put(ws, r, 1, "Column Sample Size", font=LABEL, align=CENTER, fill=BASE_FILL, border=BASE_BORDER)
        for j, (_, bc) in enumerate(cols):
            _put(ws, r, 2 + j, int(round(calc.base(table.items[0], bc))), fmt="0", fill=BASE_FILL, border=BASE_BORDER)
        r += 1
        if stat:
            _put(ws, r, 1, "Column Names", font=LABEL, align=LEFT)
            for j in range(width):
                _put(ws, r, 2 + j, string.ascii_uppercase[j], align=CENTER)
            r += 1
    _put(ws, r, 1, title if banner else f"{table.title} SUMMARY", font=TITLE, align=None, fill=None)
    footer = _sample_line(calc, table, n_cases)
    if stat:
        footer += "; " + _stat_note(calc.spec)
    _put(ws, r + 1, 1, footer, font=TITLE, align=None, fill=None)
    return r + 2 + GAP


def _pct(level: float) -> str:
    return f"{level * 100:.1f}".rstrip("0").rstrip(".") + "%"


def _stat_note(spec) -> str:
    if spec.secondary_level is None:
        symbols = f"A, B, C... (confidence level >= {_pct(spec.significance_level)})"
    else:
        symbols = (f"a, b, c... (confidence level >= {_pct(spec.secondary_level)}), "
                   f"A, B, C... (confidence level >= {_pct(spec.significance_level)})")
    note = (f"{_pct(spec.secondary_level or spec.significance_level)} confidence level; "
            f"Column comparison symbols: {symbols}; No test symbol: -")
    return note + ("" if spec.bonferroni else "; no multiple-comparison correction")


def _sheet(wb, title: str):
    ws = wb.create_sheet(title=title[:31])
    _put(ws, 1, 1, title, font=SHEET_TITLE, align=None, fill=None)
    ws.column_dimensions["A"].width = 60
    return ws


def _widths(ws, n_cols: int):
    from openpyxl.utils import get_column_letter
    for c in range(2, n_cols + 2):
        ws.column_dimensions[get_column_letter(c)].width = 13


def build_displayr_workbook(result, spec, data) -> bytes:
    calc = Calc(result, spec, data)
    tables = _tables(result, spec)
    for table in tables:
        _layout_rows(table, spec, calc)
    n_cases = len(data.df)

    wb = Workbook()
    wb.remove(wb.active)
    total_cols = [("Total", calc.total)] if calc.total else []
    if total_cols:
        ws = _sheet(wb, "Total")
        r = FIRST_ROW
        for table in tables:
            r = _write_table(ws, r, table, calc, total_cols, stat=False, banner=None, n_cases=n_cases)
        _widths(ws, max((len(t.items) for t in tables), default=1))

    groups: dict[int, list] = {}
    for bc in result.banner_columns:
        if bc.banner_var != "_total_":
            groups.setdefault(bc.banner_index, []).append(bc)
    for group in groups.values():
        banner = group[0].banner_label
        cols = [(bc.value_label, bc) for bc in group] + total_cols
        for stat in (False, True):
            ws = _sheet(wb, f"By {banner}" + (" (Stat)" if stat else ""))
            r = FIRST_ROW
            for table in tables:
                r = _write_table(ws, r, table, calc, cols, stat=stat, banner=banner, n_cases=n_cases)
            _widths(ws, len(cols) * max((len(t.items) for t in tables), default=1))

    if not wb.worksheets:
        wb.create_sheet("Tables")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
